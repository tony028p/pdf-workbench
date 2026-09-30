"""表格偵測:有框線(含合併儲存格)、無框線(目錄)、不可誤判(分欄文字、圖表格線),以及各種輸出格式。"""
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from conftest import unzip
from test_docx import all_xml_text, export_docx

RULED_TSV = 'Fruit\tSales\t\n\tQ1\tQ2\nApple\t3\t45\nBanana\t12\t60\nCherry\t100\t350'


def test_extract_ruled_table_rows(app, files):
    app.load(files['table'])
    assert app.extract(sep=False) == 'Sales Report\n\n' + RULED_TSV + '\n\nText after the table.'


def test_extract_toc_as_rows(app, files):
    app.load(files['toc'])
    text = app.extract(sep=False)
    assert 'Acknowledgements\t7\nEditorial Resilient growth\t9\n1. General assessment of the situation\t11\nIntroduction\t11' in text
    # 原本會被合成一段的兩行,現在各自一列
    assert 'The productivity slowdown has been underpinned by a decline\t68\nThe case for a regulatory reset\t69' in text


def test_two_column_text_is_not_a_table(app, files):
    app.load(files['cols'])
    text = app.extract(sep=False)
    assert '\t' not in text
    assert text.index('Left column line eight') < text.index('Right column line one')


def test_chart_grid_is_not_a_table(app, tmp_path):
    """圖表格線(大部分格子是空的)不可當成表格。"""
    p = tmp_path / 'chart.pdf'
    c = canvas.Canvas(str(p), pagesize=letter)
    for k in range(6):
        c.line(100, 400 + k * 40, 500, 400 + k * 40)
        c.line(100 + k * 80, 400, 100 + k * 80, 600)
    c.setFont('Helvetica', 9)
    for k in range(6):
        c.drawString(70, 397 + k * 40, str(k * 10))        # 軸標籤在格線外
    c.drawString(190, 470, 'peak')
    c.showPage(); c.save()
    app.load(p)
    assert '\t' not in app.extract(sep=False)


def test_table_on_rotated_page(app, files):
    app.load(files['table'])
    app.click_thumb(0)
    app.op('rotR')
    assert RULED_TSV in app.extract(sep=False)


def test_redaction_inside_table(app, files):
    app.load(files['table'])
    app.page.evaluate('setZoom(50)')
    app.tool('redact')
    # Banana 那一列:y = 792-(624..602) → 比例約 0.212–0.240
    app.drag_on_page(0, (0.10, 0.215), (0.75, 0.235))
    text = app.extract(sep=False)
    assert 'Banana' not in text and '\t12\t' not in text
    assert 'Apple\t3\t45' in text and 'Cherry\t100\t350' in text
    app.page.click('#exClose')
    data = export_docx(app)
    xml = all_xml_text(data)
    assert 'Banana' not in xml and 'Apple' in xml


def test_docx_ruled_table_with_merged_cells(app, files):
    app.load(files['table'])
    data = export_docx(app)
    doc = Document(io.BytesIO(data))
    [t] = doc.tables
    assert (len(t.rows), len(t.columns)) == (5, 3)
    assert t.cell(0, 1)._tc is t.cell(0, 2)._tc           # Sales 橫跨兩欄
    assert t.cell(0, 1).text == 'Sales' and t.cell(1, 2).text == 'Q2'
    assert [t.cell(r, 0).text for r in range(2, 5)] == ['Apple', 'Banana', 'Cherry']
    assert 'w:tblBorders' in all_xml_text(data)           # 有框線
    # 順序:標題 → 表格 → 表格後的文字
    body = [el.tag.split('}')[1] for el in doc.element.body]
    texts = [p.text for p in doc.paragraphs if p.text]
    assert body.index('tbl') > 0 and texts == ['Sales Report', 'Text after the table.']
    paras_before = [el for el in doc.element.body[:body.index('tbl')] if el.tag.endswith('}p')]
    assert any('Sales Report' in ''.join(el.itertext()) for el in paras_before)


def test_docx_toc_table(app, files):
    app.load(files['toc'])
    data = export_docx(app)
    doc = Document(io.BytesIO(data))
    [t] = doc.tables
    assert len(t.columns) == 2 and len(t.rows) == 11
    assert t.cell(0, 0).text == 'Acknowledgements' and t.cell(0, 1).text == '7'
    assert t.cell(1, 1).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT        # 頁碼靠右
    grid = [int(g.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w')) / 20 for g in t._tbl.tblGrid]
    assert grid[1] >= 40, grid                            # 頁碼欄夠寬,兩位數頁碼不會斷成兩行
    intro = t.cell(3, 0).paragraphs[0]
    assert intro.text == 'Introduction' and abs(intro.paragraph_format.left_indent.pt - 14) < 2   # 子項目縮排
    xml = all_xml_text(data)
    assert 'w:tblBorders' not in xml                      # 無框線
    assert '| 3' not in xml and '>3<' not in xml.split('<w:tbl>')[0]   # 頁首頁碼已移除
    assert doc.paragraphs[0].text == 'Table of contents' and doc.paragraphs[0].style.name == 'Heading 1'


def test_markdown_tables(app, files):
    app.load(files['table'], files['toc'])
    name, data = app.export(fmt='md')
    text = data.decode('utf-8')
    assert '| **Fruit** | **Sales** |  |\n| --- | --- | --- |\n|  | **Q1** | **Q2** |\n| Apple | 3 | 45 |' in text
    assert '| Acknowledgements | 7 |\n| --- | ---: |\n| Editorial Resilient growth | 9 |' in text


def test_html_tables(app, files):
    app.load(files['table'], files['toc'])
    _, data = app.export(fmt='html')
    doc = data.decode('utf-8')
    assert '<table class="ruled">' in doc and '<table class="plain">' in doc
    assert '<td colspan="2"><strong>Sales</strong></td>' in doc
    assert '<td style="text-align:right">7</td>' in doc


SHADED_TSV = ('Risk\tAsset\tDirection\n'
              'Middle East war\tBrent crude oil VIX U.S. high yield credit\tUp Up Down\n'
              'Energy security\tU.S. natural gas Energy equipment and services Euro\tUp Up Down\n'
              'Trade protectionism\tSpecialty retail, consumer durables and furnishings for homes Two-year Treasury\tDown Down Down')


def test_extract_table_drawn_with_filled_cells(app, files):
    """每格各自填色(多邊形)、只有橫線與一條直線的表格:填色矩形的邊也算框線(仿真實研究報告的表格)。"""
    app.load(files['shaded'])
    assert app.extract(sep=False) == 'Key scenario variables\n\n' + SHADED_TSV + '\n\nSource: test data after the table.'


def test_docx_filled_cell_table_lines_and_shading(app, files):
    app.load(files['shaded'])
    data = export_docx(app)
    [t] = Document(io.BytesIO(data)).tables
    assert len(t.rows) == 4 and len(t.columns) == 3
    # 格子裡逐行列出的項目各自成段(下一行的第一個字放得進上一行卻換行 → 刻意換行)
    assert [p.text for p in t.cell(1, 1).paragraphs] == ['Brent crude oil', 'VIX', 'U.S. high yield credit']
    assert [p.text for p in t.cell(3, 2).paragraphs] == ['Down', 'Down', 'Down']
    assert [p.text for p in t.cell(2, 1).paragraphs] == ['U.S. natural gas', 'Energy equipment and services', 'Euro']
    # 放不下下一個字而換行的是同一段
    assert [p.text for p in t.cell(3, 1).paragraphs] == ['Specialty retail, consumer durables and furnishings for homes', 'Two-year Treasury']
    W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    fill = lambda r, c: (lambda s: s.get(W + 'fill') if s is not None else None)(t.cell(r, c)._tc.find(f'{W}tcPr/{W}shd'))
    assert fill(0, 0) == 'FF4712'                                       # 表頭橘底
    assert fill(1, 1) == 'F2F2F2' and fill(2, 1) is None and fill(3, 2) == 'F2F2F2'
    # 深色底上的字用白色
    assert t.cell(0, 0).paragraphs[0].runs[0].font.color.rgb == RGBColor(0xFF, 0xFF, 0xFF)
    assert t.cell(1, 0).paragraphs[0].runs[0].font.color.rgb is None


def test_html_filled_cell_shading(app, files):
    app.load(files['shaded'])
    _, data = app.export(fmt='html')
    html = data.decode('utf-8')
    assert 'background:#F2F2F2' in html and 'color:#fff' in html
