"""匯出 Word(.docx):結構正確、內容與樣式、圖片、頁首頁尾、塗黑、實際用 LibreOffice 開啟。"""
import io
import os
import shutil
import subprocess
import zipfile
from xml.dom import minidom

import pytest
from docx import Document
from docx.shared import Pt
from PIL import Image

from conftest import pdf_texts

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def export_docx(app, **kw):
    name, data = app.export(fmt='docx', **kw)
    assert name.endswith('.docx')
    return data


def body_paras(data):
    """(樣式名稱, 文字) 的清單,略過空段落與只有圖片的段落。"""
    doc = Document(io.BytesIO(data))
    return [(p.style.name, p.text) for p in doc.paragraphs if p.text.strip()]


def all_xml_text(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return b''.join(z.read(n) for n in z.namelist() if n.endswith('.xml') or n.endswith('.rels')).decode('utf-8')


def images(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return [Image.open(io.BytesIO(z.read(n))).convert('RGB') for n in sorted(z.namelist()) if n.startswith('word/media/')]


def test_package_is_well_formed(app, files):
    app.load(files['abc'], files['img'])
    data = export_docx(app)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        assert names[0] == '[Content_Types].xml'
        for n in names:
            if n.endswith('.xml') or n.endswith('.rels'):
                minidom.parseString(z.read(n))            # 每個 XML 都要能解析
        ct = z.read('[Content_Types].xml').decode()
        rels = z.read('word/_rels/document.xml.rels').decode()
        for n in names:                                   # 每張圖片都有關聯
            if n.startswith('word/media/'):
                assert n.replace('word/', '') in rels
        assert 'wordprocessingml.document.main+xml' in ct


def test_editable_text_and_page_sections(app, files):
    app.load(files['abc'])
    data = export_docx(app)
    assert [t for _, t in body_paras(data)] == ['PAGE-A', 'PAGE-B', 'PAGE-C']
    doc = Document(io.BytesIO(data))
    assert len(doc.sections) == 3                         # 每頁一節(新頁)
    s = doc.sections[0]
    assert (round(s.page_width.pt), round(s.page_height.pt)) == (612, 792)
    # 不分頁:只有一節
    assert len(Document(io.BytesIO(export_docx(app, doc_break=False))).sections) == 1


def test_headings_bold_italic_and_order(app, files):
    app.load(files['cols'], files['img'])
    paras = body_paras(export_docx(app))
    texts = [t for _, t in paras]
    # 22pt 的「Report With Image」是 H1,20pt 的兩欄標題是 H2(依字級排層級)
    assert paras[0] == ('Heading 2', 'Two Column Title Spanning The Page')
    left = next(i for i, t in enumerate(texts) if t.startswith('Left column line one'))
    right = next(i for i, t in enumerate(texts) if t.startswith('Right column line one'))
    assert left < right
    assert ('Heading 1', 'Report With Image') in paras
    doc = Document(io.BytesIO(export_docx(app)))
    italic = [r for p in doc.paragraphs for r in p.runs if r.text == 'Text above the picture.']
    assert italic and italic[0].italic


def test_body_size_per_source_document(app, files):
    """合併內文字級不同的文件時,14pt 的中文段落不可被誤判為標題(各檔各自判斷內文字級)。"""
    app.load(files['cols'], files['paras'])
    paras = dict((t, s) for s, t in body_paras(export_docx(app)))
    assert paras['這是一段很長的中文段落，中間換行的地方不應該出現空格。'] == 'Normal'
    assert paras['Second paragraph starts here.'] == 'Normal'


def test_cjk_paragraphs_and_fonts(app, files):
    app.load(files['paras'])
    data = export_docx(app)
    assert [t for _, t in body_paras(data)] == [
        'This paragraph is a simple example of wrapped text that should be joined into one line.',
        'Second paragraph starts here.',
        '這是一段很長的中文段落，中間換行的地方不應該出現空格。',
    ]
    xml = all_xml_text(data)
    assert 'w:eastAsia="新細明體"' in xml                  # 明體 → 新細明體
    assert 'w:eastAsia="zh-TW"' in xml


def test_image_is_placed_between_paragraphs(app, files):
    app.load(files['img'])
    data = export_docx(app)
    doc = Document(io.BytesIO(data))
    seq = ['IMG' if p._p.xpath('.//pic:pic') else p.text for p in doc.paragraphs if p.text.strip() or p._p.xpath('.//pic:pic')]
    assert seq == ['Report With Image', 'Text above the picture.', 'IMG', 'Text below the picture.']
    pic = next(p for p in doc.paragraphs if p._p.xpath('.//pic:pic'))
    assert pic.paragraph_format.left_indent.pt == pytest.approx(0, abs=2)   # 原稿靠左,不置中
    assert pic.alignment is None
    [shape] = doc.inline_shapes
    assert shape.width.pt == pytest.approx(300, abs=3) and shape.height.pt == pytest.approx(150, abs=3)
    [img] = images(data)
    r, g, b = img.getpixel((img.width // 2, img.height // 2))
    assert b > 200 and r < 60                             # 藍色圖片


def test_header_footer_and_page_numbers_removed(app, files):
    app.load(files['hf'])
    data = export_docx(app)
    texts = [t for _, t in body_paras(data)]
    assert texts == ['Body text of page 1.', 'Body text of page 2.', 'Body text of page 3.']
    xml = all_xml_text(data)
    assert 'PAGE' in xml and 'footer1.xml' in xml         # 改用 Word 自動頁碼


def test_scanned_page_becomes_picture(app, files):
    app.load(files['abc'], files['jpg'])
    data = export_docx(app, scope='all')
    doc = Document(io.BytesIO(data))
    assert len(doc.inline_shapes) == 1
    [img] = images(data)
    assert img.getpixel((img.width // 2, img.height // 2))[0] > 180


def test_redacted_text_and_image_region_excluded(app, files):
    """塗黑:文字不可出現在任何 XML;圖片被塗黑的部分在裁切出來的圖裡是黑的。"""
    app.load(files['secret'], files['img'])
    app.tool('redact')
    app.drag_on_page(0, (0.05, 0.10), (0.80, 0.25))       # 蓋住 SECRET-123
    app.page.evaluate('setZoom(33)')
    app.click_thumb(2)
    app.drag_on_page(2, (0.10, 0.25), (0.35, 0.45))        # 蓋住圖片左半邊
    data = export_docx(app)
    xml = all_xml_text(data)
    assert 'SECRET' not in xml
    assert 'KEEP-456' in xml
    [img] = images(data)
    assert max(img.getpixel((int(img.width * 0.1), img.height // 2))) < 60      # 左邊被塗黑
    assert img.getpixel((int(img.width * 0.9), img.height // 2))[2] > 200       # 右邊仍是藍色


def test_image_mode_one_picture_per_page_no_text(app, files):
    app.load(files['secret'])
    app.tool('redact')
    app.drag_on_page(0, (0.05, 0.10), (0.80, 0.25))
    data = export_docx(app, doc_mode='image')
    doc = Document(io.BytesIO(data))
    assert len(doc.inline_shapes) == 2 and len(doc.sections) == 2
    assert 'SECRET' not in all_xml_text(data) and 'KEEP' not in all_xml_text(data)
    top = images(data)[0]
    assert max(top.getpixel((int(top.width * 0.3), int(top.height * 0.17)))) < 60   # 塗黑有燒進圖


def test_rotated_page_orientation(app, files):
    app.load(files['abc'])
    app.click_thumb(0)
    app.op('rotR')
    doc = Document(io.BytesIO(export_docx(app)))
    s = doc.sections[0]
    assert s.page_width > s.page_height and s.orientation == 1   # 橫向
    assert body_paras(export_docx(app))[0][1] == 'PAGE-A'


def test_selected_pages_only(app, files):
    app.load(files['abc'])
    app.click_thumb(1)
    assert [t for _, t in body_paras(export_docx(app, scope='sel'))] == ['PAGE-B']


@pytest.mark.skipif(not shutil.which('soffice') and not os.environ.get('REQUIRE_SOFFICE'), reason='需要 LibreOffice(含 Writer)')
def test_opens_in_libreoffice(app, files, tmp_path):
    """用 LibreOffice 實際開啟並轉成 PDF:能開、頁數正確、文字都在。"""
    app.load(files['cols'], files['paras'], files['img'], files['hf'], files['table'], files['toc'])
    for mode in ('edit', 'image'):
        src = tmp_path / f'{mode}.docx'
        src.write_bytes(export_docx(app, doc_mode=mode))
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', str(tmp_path), str(src)],
                       check=True, capture_output=True, timeout=180)
        pdf = (tmp_path / f'{mode}.pdf').read_bytes()
        texts = pdf_texts(pdf)
        assert len(texts) == 8, (mode, len(texts))         # 6 個檔共 8 頁,每頁一節
        if mode == 'edit':
            joined = '\n'.join(texts)
            for s in ('Two Column Title', 'example of wrapped text', '這是一段很長的中文段落', 'Report With Image', 'Body text of page 3.',
                      'Banana', 'Recent Developments'):
                assert s in joined, s
