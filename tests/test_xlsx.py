"""表格匯出 Excel(.xlsx):工作表結構、合併儲存格、數字存成數值並保留顯示格式(千分位、小數、百分比、負數、貨幣)、
底色與粗體、格內換行、塗黑排除;用 openpyxl 讀取,並用 LibreOffice 實際開啟、轉成 CSV 比對顯示文字。"""
import csv
import io
import os
import shutil
import subprocess
import zipfile
from xml.dom import minidom

import openpyxl
import pytest

from fixtures import OCR_TABLE


def export_xlsx(app):
    name, data = app.export(fmt='xlsx')
    assert name.endswith('.xlsx')
    return data


def workbook(data):
    return openpyxl.load_workbook(io.BytesIO(data))


def values(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_package_well_formed_and_ruled_table(app, files):
    app.load(files['table'])
    data = export_xlsx(app)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for n in z.namelist():
            minidom.parseString(z.read(n))
    wb = workbook(data)
    assert wb.sheetnames == ['第1頁']
    ws = wb['第1頁']
    assert values(ws) == [['Fruit', 'Sales', None], [None, 'Q1', 'Q2'], ['Apple', 3, 45], ['Banana', 12, 60], ['Cherry', 100, 350]]
    assert 'B1:C1' in [str(r) for r in ws.merged_cells.ranges]           # Sales 橫跨兩欄
    assert ws['A1'].font.b and not ws['A3'].font.b
    assert ws['B3'].alignment.horizontal == 'right'


def test_numbers_become_values_with_original_format(ocr_app, files):
    """財報式數字表(辨識掃描檔):每個數字都是數值,而且格式顯示出來與原文相同。"""
    ocr_app.load(files['scan_numbers'])
    ocr_app.ocr('auto')
    ws = workbook(export_xlsx(ocr_app))['第1頁']
    cells = {row[0].value: row for row in ws.iter_rows(min_row=2)}
    expect = {
        '營業收入': (12345678.9, '#,##0.00'), '營業成本': (-4321000.5, '#,##0.00;(#,##0.00)'),
        '研發支出': (1000000, '#,##0'), '利息收入': (0.25, '0.00'), '匯率': (31.4567, '0.0000'),
        '成長率': (0.158, '+0.0%;-0.0%;0'), '每股淨值': (2499.99, '"NT$ "#,##0.00'), '總資產': (1234567890.12, '#,##0.00'),
    }
    for label, (v, fmt) in expect.items():
        c = cells[label][1]
        assert isinstance(c.value, (int, float)) and abs(c.value - v) < 1e-6 * max(1, abs(v)), (label, c.value)
        assert c.number_format == fmt, (label, c.number_format)
    assert cells['成長率'][2].value == pytest.approx(-0.0235) and cells['成長率'][2].number_format == '0.00%'
    assert cells['總資產'][2].value == pytest.approx(100000.001) and cells['總資產'][2].number_format == '#,##0.000'


def libreoffice_csv(data, tmp_path):
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        if os.environ.get('REQUIRE_SOFFICE'):
            pytest.fail('CI 必須安裝 LibreOffice')
        pytest.skip('沒有 LibreOffice')
    src = tmp_path / 'table.xlsx'
    src.write_bytes(data)
    subprocess.run([soffice, '--headless', '--convert-to', 'csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,true', '--outdir', str(tmp_path), str(src)],
                   check=True, capture_output=True, timeout=180)
    return list(csv.reader(io.StringIO((tmp_path / 'table.csv').read_text(encoding='utf-8'))))


def test_opens_in_libreoffice_and_displays_like_original(ocr_app, files, tmp_path):
    """LibreOffice 開啟後以「顯示的文字」轉成 CSV:每一格都要和原本表格上的字一模一樣。"""
    ocr_app.load(files['scan_numbers'])
    ocr_app.ocr('auto')
    rows = libreoffice_csv(export_xlsx(ocr_app), tmp_path)
    assert rows == [list(r) for r in OCR_TABLE], rows


def test_shaded_table_fill_bold_and_line_breaks(app, files):
    app.load(files['shaded'])
    ws = workbook(export_xlsx(app))['第1頁']
    assert ws['A1'].value == 'Risk' and ws['A1'].font.b
    assert ws['A1'].fill.fgColor.rgb == 'FFFF4712' and ws['A1'].font.color.rgb == 'FFFFFFFF'    # 深色底白字
    assert ws['B2'].fill.fgColor.rgb == 'FFF2F2F2'
    assert ws['B2'].value == 'Brent crude oil\nVIX\nU.S. high yield credit' and ws['B2'].alignment.wrap_text
    assert ws['C4'].value == 'Down\nDown\nDown'


def test_toc_and_multiple_tables(app, files):
    app.load(files['toc'], files['table'])
    wb = workbook(export_xlsx(app))
    assert len(wb.sheetnames) == 2 and wb.sheetnames[1] == '第2頁'
    toc = values(wb[wb.sheetnames[0]])
    assert ['Acknowledgements', 7] in toc                                 # 頁碼是數值


def test_redaction_excluded(app, files):
    app.load(files['table'])
    app.page.evaluate('setZoom(50)')
    app.tool('redact')
    app.drag_on_page(0, (0.10, 0.215), (0.75, 0.235))                     # Banana 那一列
    data = export_xlsx(app)
    flat = [v for row in values(workbook(data).active) for v in row]
    assert 'Banana' not in flat and 12 not in flat and 60 not in flat
    assert 'Apple' in flat and 350 in flat
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert not any(b'Banana' in z.read(n) for n in z.namelist())


def test_no_tables_message(app, files):
    app.load(files['paras'])
    p = app.page
    p.click('#btnExport')
    p.check('input[name=xfmt][value=xlsx]')
    p.click('#xOk')
    app.wait_idle()
    assert '沒有偵測到表格' in app.toast()


def test_number_parsing_rules(app):
    """文字→數字的規則:代號(開頭是 0)、日期、電話維持文字。"""
    r = app.page.evaluate("""['1,234', '-12.50', '(3,000)', '12.5%', '+1.2%', 'US$ 9.80', '$5', '00123', '2026/09/30', '2735-8861', '1,23', 'abc', '3.']
        .map(s => { const n = parseNumberCell(s); return n ? [n.value, n.fmt] : null; })""")
    assert r == [[1234, '#,##0'], [-12.5, '0.00'], [-3000, '#,##0;(#,##0)'], [0.125, '0.0%'], [0.012, '+0.0%;-0.0%;0'],
                 [9.8, '"US$ "0.00'], [5, '"$"0'], None, None, None, None, None, None]
